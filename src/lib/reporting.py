import pandas as pd
import matplotlib.pyplot as plt
from openpyxl import load_workbook
from openpyxl.chart import LineChart, Reference

plt.rcParams["font.sans-serif"] = ["Microsoft JhengHei"]
plt.rcParams["axes.unicode_minus"] = False

# 輸出至 Excel（含圖表）
def out_excel(name, df, result):
    writer = pd.ExcelWriter(name + ".xlsx", engine="openpyxl")
    df.to_excel(writer, sheet_name="Sheet1", index=False)
    result.to_excel(writer, sheet_name="result", index=False)
    writer.close()

    wb = load_workbook(name + ".xlsx")
    ws = wb["result"]
    chart = LineChart()
    chart.title = "策略累積報酬"
    chart.style = 2
    chart.y_axis.title = "累積損益"
    chart.x_axis.title = "資料點"

    # 假設 'cus' 在 Sheet1 的第7欄（G），列索引從2開始
    data = Reference(
        wb["Sheet1"],
        min_col=df.columns.get_loc("cus") + 1,
        min_row=2,
        max_row=len(df) + 1,
    )
    chart.add_data(data, titles_from_data=False)
    ws.add_chart(chart, "H2")

    wb.save(name + ".xlsx")
